from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'TechScape_India_Project_Report.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.85); sec.bottom_margin = Inches(.85)
sec.left_margin = Inches(.9); sec.right_margin = Inches(.9)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.space_after = Pt(7)
for name, size in [('Title', 26), ('Heading 1', 16), ('Heading 2', 12)]:
    styles[name].font.name = 'Aptos Display'; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor(0,0,0)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def heading(text): doc.add_heading(text, level=1)
def bullet(text): doc.add_paragraph(text, style='List Bullet')

p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('TechScape India')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('An AR Based Interactive Visualization of India’s Technology and Innovation Ecosystem'); r.bold=True; r.font.size=Pt(14)
doc.add_paragraph('Project Report', style='Subtitle').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('\n\n')
table = doc.add_table(rows=4, cols=2); table.style='Table Grid'
for i,(a,b) in enumerate([('Submitted by','[Add group member names and roll numbers]'),('Course','[Add course name]'),('Institution','[Add institution name]'),('Academic Year','[Add academic year]')]):
    table.cell(i,0).text=a; table.cell(i,1).text=b; shade(table.cell(i,0),'D9EAF7')
doc.add_page_break()

heading('Group Members')
doc.add_paragraph('Please complete the following details before submission.')
t=doc.add_table(rows=1, cols=3); t.style='Table Grid'
for c,text in zip(t.rows[0].cells,['No.','Member Name','Roll Number']): c.text=text; shade(c,'1F4E78');
for c in t.rows[0].cells:
    for run in c.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True
for i in range(1,5):
    row=t.add_row().cells; row[0].text=str(i); row[1].text=''; row[2].text=''

heading('Introduction')
doc.add_paragraph('TechScape India is an augmented reality application that presents India’s technology and innovation ecosystem through an interactive AR map. Instead of a conventional static map, users scan a TechScape India image target to view a map and explore technology information for selected Indian states.')
doc.add_paragraph('The application currently focuses on Karnataka, Maharashtra, Telangana, Tamil Nadu, Kerala, Gujarat, Uttar Pradesh, West Bengal and Rajasthan. Each supported state presents concise information about its technology ecosystem, major cities, key sectors, companies and startups, research institutions, initiatives and an interesting technology fact.')

heading('Requirements')
bullet('Unity 2022 LTS with Android build support')
bullet('Vuforia Engine and an active image target database')
bullet('Android smartphone with camera permission')
bullet('TechScape India printed or on-screen scan target')
bullet('ARM64 Android device for APK installation')
bullet('Structured JSON dataset for state technology information')

heading('Implementation')
doc.add_paragraph('The project was developed in Unity using Vuforia image-target tracking. When the TechScape India target is detected, the application displays an AR India map. State marker selection is restricted to the nine supported TechScape states, while the full India map geometry remains visible.')
doc.add_paragraph('State information is maintained in a JSON resource file and is injected into the existing state-detail interface at runtime. The detail screen follows a consistent structure: Tech Overview, City Spotlights, Key Sectors, Companies and Startups, Innovation and Research, Initiatives and Tech Fact. The information panel uses vertical touch scrolling so that users can read the complete content on a mobile device.')
doc.add_paragraph('Navigation uses an explicit list of the nine supported states rather than Unity build-index order. This prevents the Next control from opening legacy art and culture scenes. Unsupported marker objects are disabled at runtime, removing their thumbnails, colliders and click targets from the AR map.')

heading('Results')
doc.add_paragraph('The final application provides a focused technology-based AR experience for nine Indian state ecosystems. The application title and Android package identify it as TechScape India, and the Vuforia target has been configured specifically for the project.')
for s in ['Karnataka: Bengaluru and Mysuru','Maharashtra: Mumbai, Pune and Navi Mumbai','Telangana: Hyderabad and Warangal','Tamil Nadu: Chennai, Coimbatore and Hosur','Kerala: Kochi, Thiruvananthapuram and Kozhikode','Gujarat: Ahmedabad, Gandhinagar and Surat','Uttar Pradesh: Noida, Greater Noida and Lucknow','West Bengal: Kolkata, Salt Lake Sector V and New Town','Rajasthan: Jaipur, Kota and Udaipur']:
    bullet(s)
doc.add_paragraph('The Android APK was built for ARM64 devices. The completed app displays only the supported TechScape markers, presents technology-focused information through scrollable detail panels, and keeps navigation within the supported state set.')

doc.save(OUT)
