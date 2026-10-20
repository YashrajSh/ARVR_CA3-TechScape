from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = "TechScape_India_Project_Report.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_text(cell, text, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    if color:
        run.font.color.rgb = RGBColor(*color)


doc = Document(SOURCE)

# Complete the existing member list without changing the report's original structure.
members = ["Yashraj", "Aarvee", "Sara", "Tanishka"]
member_table = doc.tables[1]
for index, member in enumerate(members, start=1):
    member_table.cell(index, 1).text = member

# Add a concise, scan-friendly contribution matrix directly below the group list.
heading = doc.add_paragraph("Group Contributions", style="Heading 1")
heading.paragraph_format.space_before = Pt(12)
heading.paragraph_format.space_after = Pt(6)

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.autofit = False
table.columns[0].width = Inches(1.35)
table.columns[1].width = Inches(5.45)

headers = ["Member", "Contribution"]
for i, header in enumerate(headers):
    cell = table.cell(0, i)
    set_cell_shading(cell, "17365D")
    set_cell_border(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_text(cell, header, bold=True, color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

contributions = [
    ("Yashraj", "Project coordination, Unity and Vuforia AR integration, image-target configuration, and Android APK build preparation."),
    ("Aarvee", "Technology research and structuring of the state and city technology ecosystem content."),
    ("Sara", "AR user-interface planning, scrollable information-panel design, and visual content review."),
    ("Tanishka", "Functional testing, state-navigation validation, documentation, and project presentation preparation."),
]

for row_index, (member, contribution) in enumerate(contributions, start=1):
    cells = table.add_row().cells
    for col_index, value in enumerate((member, contribution)):
        cell = cells[col_index]
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_index % 2 == 0:
            set_cell_shading(cell, "F2F6FA")
        set_text(cell, value, bold=(col_index == 0), align=WD_ALIGN_PARAGRAPH.LEFT)

# Move the newly-created heading and table from the document end to immediately after the member table.
member_table._tbl.addnext(heading._p)
heading._p.addnext(table._tbl)

doc.save(SOURCE)
